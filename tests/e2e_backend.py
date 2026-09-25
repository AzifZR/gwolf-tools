"""E2E test: run against live app.py server on 127.0.0.1:8083."""
import io
import json
import random
import sys
import urllib.request
import urllib.error
import uuid
import zipfile
from pathlib import Path
from PIL import Image, ImageDraw
import pypdf
from docx import Document

TMP = Path(__file__).parent / "fixtures"
TMP.mkdir(exist_ok=True)
BASE = "http://127.0.0.1:8083"

failures = []

def check(name, cond, detail=""):
    status = "PASS" if cond else "FAIL"
    print(f"[{status}] {name} {detail}")
    if not cond:
        failures.append(name)

def multipart(fields, files):
    boundary = uuid.uuid4().hex
    out = io.BytesIO()
    for k, v in fields.items():
        out.write(f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"\r\n\r\n{v}\r\n".encode())
    for field, fn, data, ct in files:
        out.write(f"--{boundary}\r\nContent-Disposition: form-data; name=\"{field}\"; filename=\"{fn}\"\r\n"
                  f"Content-Type: {ct}\r\n\r\n".encode())
        out.write(data)
        out.write(b"\r\n")
    out.write(f"--{boundary}--\r\n".encode())
    return out.getvalue(), f"multipart/form-data; boundary={boundary}"

def post(path, fields=None, files=None):
    body, ct = multipart(fields or {}, files or [])
    req = urllib.request.Request(BASE + path, data=body, headers={"Content-Type": ct}, method="POST")
    try:
        with urllib.request.urlopen(req, timeout=120) as r:
            return r.status, {k.lower(): v for k, v in r.headers.items()}, r.read()
    except urllib.error.HTTPError as e:
        return e.code, {k.lower(): v for k, v in e.headers.items()}, e.read()
    except Exception as e:
        return 500, {}, str(e).encode()

def get(path):
    req = urllib.request.Request(BASE + path)
    with urllib.request.urlopen(req, timeout=30) as r:
        return r.status, {k.lower(): v for k, v in r.headers.items()}, r.read()

def make_noisy_jpeg(path, w=400, h=300):
    rnd = random.Random(1)
    im = Image.new("RGB", (w, h))
    im.putdata([(rnd.randrange(256), rnd.randrange(256), rnd.randrange(256)) for _ in range(w * h)])
    im.save(path, "JPEG", quality=92)
    return path

def make_png(path, w=300, h=200):
    rnd = random.Random(3)
    im = Image.new("RGB", (w, h))
    im.putdata([(rnd.randrange(64), rnd.randrange(128), rnd.randrange(96)) for _ in range(w * h)])
    im.save(path, "PNG", optimize=True)
    return path

def make_pdf(path, pages=2):
    imgs = []
    for p in range(pages):
        im = Image.new("RGB", (400, 600), (255, 255, 255))
        d = ImageDraw.Draw(im)
        d.text((50, 50), f"PAGE {p+1} TEST", fill=(0, 0, 0))
        imgs.append(im)
    imgs[0].save(path, "PDF", save_all=True, append_images=imgs[1:], resolution=100.0)
    return path

def make_text_pdf(path, pages=2):
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        c = canvas.Canvas(str(path), pagesize=A4)
        for p in range(pages):
            c.setFont("Helvetica-Bold", 16)
            c.drawString(60, 760, f"WORDTEST PAGE {p+1} MARKER-ALPHA")
            c.setFont("Helvetica", 12)
            c.drawString(60, 720, f"Paragraf uji halaman {p+1} semangka stroberi.")
            c.showPage()
        c.save()
        return path
    except Exception:
        return make_pdf(path, pages)

def make_docx(path):
    doc = Document()
    doc.add_heading("Dokumen Uji Gwolf", level=1)
    doc.add_paragraph("Baris pertama paragraf uji.")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "A1"
    t.cell(0, 1).text = "B1"
    t.cell(1, 0).text = "A2"
    t.cell(1, 1).text = "B2"
    doc.save(path)
    return path

# 1. GET / and /api/status
def test_static_and_status():
    s, h, body = get("/")
    check("GET / serves html", s == 200 and b"<html" in body.lower())
    s, h, body = get("/api/status")
    check("GET /api/status 200", s == 200)
    if s == 200:
        feats = json.loads(body)["features"]
        check("status has all 9 features", len(feats) == 9, str(feats))

# 2. PDF merge & compress
def test_pdf_tools():
    p1 = make_pdf(TMP / "doc1.pdf", pages=2).read_bytes()
    p2 = make_pdf(TMP / "doc2.pdf", pages=1).read_bytes()
    
    # Merge
    s, h, body = post("/api/pdf/merge", files=[("files", "doc1.pdf", p1, "application/pdf"), ("files", "doc2.pdf", p2, "application/pdf")])
    check("pdf/merge 200", s == 200)
    if s == 200:
        check("pdf/merge original size header", h.get("x-original-size") == str(len(p1) + len(p2)))
        rdr = pypdf.PdfReader(io.BytesIO(body))
        check("pdf/merge 3 pages", len(rdr.pages) == 3)
        
    # Compress
    s, h, body = post("/api/pdf/compress", files=[("files", "doc1.pdf", p1, "application/pdf")])
    check("pdf/compress 200", s == 200)
    if s == 200:
        check("pdf/compress headers", h.get("x-original-size") == str(len(p1)) and h.get("x-result-size") == str(len(body)))

# 3. PDF to Img
def test_pdf_to_img():
    p3 = make_pdf(TMP / "multi.pdf", pages=3).read_bytes()
    s, h, body = post("/api/pdf/to-img", fields={"format": "PNG", "dpi": "96"},
                      files=[("files", "multi.pdf", p3, "application/pdf")])
    check("pdf/to-img multi 200", s == 200)
    if s == 200:
        check("pdf/to-img X-Page-Count=3", h.get("x-page-count") == "3")
        z = zipfile.ZipFile(io.BytesIO(body))
        check("pdf/to-img zip 3 entries", len(z.namelist()) == 3)

# 4. Image operations
def test_image_tools():
    jpg = make_noisy_jpeg(TMP / "noise.jpg").read_bytes()
    png = make_png(TMP / "graphic.png").read_bytes()
    
    # img/compress
    s, h, body = post("/api/img/compress", fields={"quality": "50"}, files=[("files", "noise.jpg", jpg, "image/jpeg")])
    check("img/compress 200", s == 200)
    
    # img/upscale
    s, h, body = post("/api/img/upscale", fields={"scale": "2"}, files=[("files", "graphic.png", png, "image/png")])
    check("img/upscale 200", s == 200)
    if s == 200:
        im = Image.open(io.BytesIO(body))
        check("img/upscale dimensions 2x", im.size == (600, 400))
        
    # img/convert
    s, h, body = post("/api/img/convert", fields={"format": "WEBP"}, files=[("files", "noise.jpg", jpg, "image/jpeg")])
    check("img/convert 200", s == 200)
    if s == 200:
        im = Image.open(io.BytesIO(body))
        check("img/convert format WEBP", im.format == "WEBP")
        
    # img/to-pdf
    s, h, body = post("/api/img/to-pdf", files=[("files", "1.jpg", jpg, "image/jpeg"), ("files", "2.jpg", jpg, "image/jpeg")])
    check("img/to-pdf 200", s == 200)
    if s == 200:
        rdr = pypdf.PdfReader(io.BytesIO(body))
        check("img/to-pdf 2 pages", len(rdr.pages) == 2)

# 5. Word <-> PDF
def test_docx_tools():
    docx_b = make_docx(TMP / "uji.docx").read_bytes()
    pdf_b = make_text_pdf(TMP / "wordtest.pdf", pages=2).read_bytes()
    
    # word/to-pdf
    s, h, body = post("/api/word/to-pdf", files=[("files", "uji.docx", docx_b, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")])
    check("word/to-pdf 200", s == 200)
    if s == 200:
        check("word/to-pdf headers", h.get("x-original-size") == str(len(docx_b)) and h.get("x-result-size") == str(len(body)))
        rdr = pypdf.PdfReader(io.BytesIO(body))
        check("word/to-pdf valid >=1 page", len(rdr.pages) >= 1)
        
    # pdf/to-word
    s, h, body = post("/api/pdf/to-word", files=[("files", "wordtest.pdf", pdf_b, "application/pdf")])
    check("pdf/to-word 200", s == 200)
    if s == 200:
        check("pdf/to-word headers", h.get("x-original-size") == str(len(pdf_b)) and h.get("x-result-size") == str(len(body)))
        doc = Document(io.BytesIO(body))
        check("pdf/to-word non-empty document", len(doc.paragraphs) > 0)

# 6. Error handling
def test_errors():
    s, h, body = post("/api/word/to-pdf", files=[("files", "old.doc", b"\xd0\xcf\x11\xe0fake", "application/msword")])
    check("legacy .doc error status >=400", s >= 400 and b"error" in body)
    
    s, h, body = post("/api/word/to-pdf", files=[("files", "bad.docx", b"junk", "application/octet-stream")])
    check("garbage file error status >=400", s >= 400 and b"error" in body)
    
    s, h, body = post("/api/pdf/to-word")
    check("no file error status 400", s == 400 and b"error" in body)

if __name__ == "__main__":
    print("=== RUNNING FULL E2E BACKEND TESTS ===")
    test_static_and_status()
    test_pdf_tools()
    test_pdf_to_img()
    test_image_tools()
    test_docx_tools()
    test_errors()
    print(f"\nTOTAL FAILURES: {len(failures)}")
    if failures:
        sys.exit(1)
