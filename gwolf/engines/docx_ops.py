import os
import io
import shutil
import subprocess
import tempfile
from pathlib import Path
import pypdf
from gwolf.config import TIMEOUT_SOFFICE

def word_to_pdf_soffice(inp_bytes, filename):
    soff = shutil.which("soffice")
    if not soff:
        return None
    tmpdir = None
    try:
        tmpdir = tempfile.mkdtemp()
        suffix = Path(filename).suffix or ".docx"
        safe = "input" + suffix
        ip = os.path.join(tmpdir, safe)
        Path(ip).write_bytes(inp_bytes)
        cmd = [soff, "--headless", "--convert-to", "pdf", "--outdir", tmpdir, ip]
        subprocess.run(cmd, timeout=TIMEOUT_SOFFICE, check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        pdfs = list(Path(tmpdir).glob("*.pdf"))
        if not pdfs:
            exp = Path(tmpdir) / (Path(safe).stem + ".pdf")
            if exp.exists():
                pdfs = [exp]
            else:
                return None
        pp = max(pdfs, key=lambda p: p.stat().st_size) if len(pdfs) > 1 else pdfs[0]
        data = pp.read_bytes()
        if not data or len(data) < 80 or not data.startswith(b"%PDF"):
            return None
        try:
            r = pypdf.PdfReader(io.BytesIO(data))
            if len(r.pages) == 0:
                return None
        except Exception:
            pass
        return data
    except Exception:
        return None
    finally:
        if tmpdir:
            try:
                shutil.rmtree(tmpdir)
            except Exception:
                pass

def word_to_pdf_docx_pillow(inp_bytes):
    try:
        import docx
    except Exception:
        return None
    try:
        from PIL import Image, ImageDraw, ImageFont
    except Exception:
        return None
    try:
        doc = docx.Document(io.BytesIO(inp_bytes))
    except Exception:
        return None
    W, H = 1240, 1754
    margin = 60
    usable = W - 2 * margin

    def _get_font(sz, bold=False):
        cands = []
        if bold:
            cands += ["DejaVuSans-Bold.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", "C:\\Windows\\Fonts\\DejaVuSans-Bold.ttf", "C:\\Windows\\Fonts\\DejaVuSans.ttf"]
        cands += ["DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "C:\\Windows\\Fonts\\DejaVuSans.ttf", "/system/fonts/DroidSans.ttf"]
        for c in cands:
            try:
                return ImageFont.truetype(c, sz)
            except Exception:
                continue
        try:
            return ImageFont.load_default(size=sz)
        except TypeError:
            return ImageFont.load_default()

    font_normal = _get_font(22)
    font_h1 = _get_font(30, bold=True)
    font_h2 = _get_font(26, bold=True)
    font_h3 = _get_font(23, bold=True)
    font_tbl = _get_font(18)

    def wrap(text, font, draw):
        if not text:
            return []
        words = text.split()
        if not words:
            return [text]
        lines, cur = [], ""
        for w in words:
            tst = cur + (" " if cur else "") + w
            try:
                wd = draw.textlength(tst, font=font)
            except Exception:
                try:
                    wd = font.getlength(tst)
                except Exception:
                    wd = len(tst) * 10
            if wd <= usable:
                cur = tst
            else:
                if cur:
                    lines.append(cur)
                # hard-break long word
                if len(w) > 35:
                    chunk = ""
                    for ch in w:
                        tc = chunk + ch
                        try:
                            cwd = draw.textlength(tc, font=font)
                        except Exception:
                            cwd = len(tc) * 10
                        if cwd > usable:
                            if chunk:
                                lines.append(chunk)
                            chunk = ch
                        else:
                            chunk = tc
                    cur = chunk
                else:
                    cur = w
        if cur:
            lines.append(cur)
        return lines

    pages = []
    im = Image.new("RGB", (W, H), "white")
    draw = ImageDraw.Draw(im)
    y = margin
    line_h_normal = 32
    line_h_h1 = 42
    line_h_tbl = 26

    # map elements to preserve order
    para_map = {p._element: p for p in doc.paragraphs}
    tbl_map = {t._element: t for t in doc.tables}
    try:
        body_children = list(doc.element.body.iterchildren())
    except Exception:
        body_children = list(para_map.keys()) + list(tbl_map.keys())

    def new_page():
        nonlocal im, draw, y
        pages.append(im)
        im = Image.new("RGB", (W, H), "white")
        draw = ImageDraw.Draw(im)
        y = margin

    def ensure(h):
        nonlocal y, im, draw
        if y + h > H - margin:
            new_page()

    for child in body_children:
        try:
            if child in para_map:
                p = para_map[child]
                txt = p.text.strip() if p.text else ""
                if not txt:
                    ensure(line_h_normal // 2)
                    y += line_h_normal // 2
                    continue
                style = ""
                try:
                    style = p.style.name if p.style and p.style.name else ""
                except Exception:
                    style = ""
                if style.startswith("Heading 1") or style == "Title":
                    f, lh = font_h1, line_h_h1
                elif style.startswith("Heading 2"):
                    f, lh = font_h2, 36
                elif style.startswith("Heading 3"):
                    f, lh = font_h3, 33
                else:
                    f, lh = font_normal, line_h_normal
                ls = wrap(txt, f, draw)
                for ln in ls:
                    ensure(lh)
                    try:
                        draw.text((margin, y), ln, fill="black", font=f)
                    except Exception:
                        draw.text((margin, y), ln, fill="black")
                    y += lh
                y += 6
            elif child in tbl_map:
                t = tbl_map[child]
                for row in t.rows:
                    try:
                        cells = []
                        for c in row.cells:
                            try:
                                ct = c.text.strip().replace("\n", " ") if c.text else ""
                            except Exception:
                                ct = ""
                            cells.append(ct)
                        line = " | ".join(cells).strip()
                        if not line:
                            continue
                        ls = wrap(line, font_tbl, draw)
                        for ln in ls:
                            ensure(line_h_tbl)
                            try:
                                draw.text((margin, y), ln, fill="black", font=font_tbl)
                            except Exception:
                                draw.text((margin, y), ln, fill="black")
                            y += line_h_tbl
                    except Exception:
                        continue
                y += 8
            else:
                continue
        except Exception:
            continue

    if not pages and y == margin:
        for p in doc.paragraphs:
            try:
                txt = p.text.strip() if p.text else ""
                if not txt:
                    continue
                ls = wrap(txt, font_normal, draw)
                for ln in ls:
                    ensure(line_h_normal)
                    draw.text((margin, y), ln, fill="black", font=font_normal)
                    y += line_h_normal
                y += 4
            except Exception:
                continue
        for t in doc.tables:
            for row in t.rows:
                try:
                    line = " | ".join(c.text.strip().replace("\n", " ") for c in row.cells)
                    if not line.strip():
                        continue
                    ls = wrap(line, font_tbl, draw)
                    for ln in ls:
                        ensure(line_h_tbl)
                        draw.text((margin, y), ln, fill="black", font=font_tbl)
                        y += line_h_tbl
                except Exception:
                    continue

    pages.append(im)
    try:
        buf = io.BytesIO()
        pages[0].save(buf, "PDF", save_all=True, append_images=pages[1:], resolution=150.0)
        out = buf.getvalue()
        if not out or len(out) < 200:
            return None
        try:
            r = pypdf.PdfReader(io.BytesIO(out))
            if len(r.pages) == 0:
                return None
        except Exception:
            if not out.startswith(b"%PDF"):
                return None
        return out
    except Exception:
        return None

def pdf_to_word_pdf2docx(inp_bytes):
    try:
        import importlib
        pdf2docx_mod = importlib.import_module("pdf2docx")
        Converter = getattr(pdf2docx_mod, "Converter")
    except Exception:
        return None
    tmp_in = tmp_out = None
    try:
        fd1, p1 = tempfile.mkstemp(suffix=".pdf")
        os.close(fd1)
        fd2, p2 = tempfile.mkstemp(suffix=".docx")
        os.close(fd2)
        tmp_in, tmp_out = p1, p2
        Path(p1).write_bytes(inp_bytes)
        cv = Converter(p1)
        try:
            cv.convert(p2, start=0, end=None)
        finally:
            try:
                cv.close()
            except Exception:
                pass
        data = Path(p2).read_bytes()
        if not data or len(data) < 100:
            return None
        return data
    except Exception:
        return None
    finally:
        for p in (tmp_in, tmp_out):
            if p:
                try:
                    os.remove(p)
                except Exception:
                    pass

def pdf_to_word_pypdf_docx(inp_bytes):
    try:
        import docx
    except Exception:
        return None
    try:
        reader = pypdf.PdfReader(io.BytesIO(inp_bytes))
        n = len(reader.pages)
        if n == 0:
            return None
    except Exception:
        return None
    try:
        doc = docx.Document()
        for idx, page in enumerate(reader.pages, start=1):
            try:
                txt = page.extract_text() or ""
            except Exception:
                txt = ""
            try:
                doc.add_heading(f"Halaman {idx}", level=1)
            except Exception:
                pa = doc.add_paragraph()
                try:
                    pa.add_run(f"Halaman {idx}").bold = True
                except Exception:
                    pa.add_run(f"Halaman {idx}")
            if txt.strip():
                paras = []
                cur = []
                for line in txt.splitlines():
                    if line.strip() == "":
                        if cur:
                            paras.append("\n".join(cur))
                            cur = []
                    else:
                        cur.append(line)
                if cur:
                    paras.append("\n".join(cur))
                if not paras:
                    paras = [txt.strip()]
                for pt in paras:
                    try:
                        if pt.strip():
                            doc.add_paragraph(pt.strip())
                    except Exception:
                        continue
        buf = io.BytesIO()
        doc.save(buf)
        out = buf.getvalue()
        if not out or len(out) < 80:
            return None
        return out
    except Exception:
        return None
