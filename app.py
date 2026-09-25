#!/usr/bin/env python3
"""
Gwolf Toolbox - Local Web Utilities (PDF & Image Tools)
Running on Termux Android (Port 8083)
"""

import os
import io
import json
import uuid
import zipfile
import shutil
import subprocess
import tempfile
from pathlib import Path
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse, parse_qs
from PIL import Image, ImageOps, ImageEnhance
import pypdf

PORT = 8083
BASE_DIR = Path.home() / "tools-web"
STATIC_DIR = BASE_DIR / "static"
UPLOAD_DIR = BASE_DIR / "tmp"

UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
STATIC_DIR.mkdir(parents=True, exist_ok=True)


def _gs_compress(input_bytes):
    gs = shutil.which("gs")
    if not gs:
        for cand in ("gswin64c", "gswin32c", "gsc"):
            gs = shutil.which(cand)
            if gs:
                break
    if not gs:
        return None
    tmp_in = tmp_out = None
    try:
        fd_in, p_in = tempfile.mkstemp(suffix=".pdf")
        os.close(fd_in)
        fd_out, p_out = tempfile.mkstemp(suffix=".pdf")
        os.close(fd_out)
        tmp_in, tmp_out = p_in, p_out
        Path(p_in).write_bytes(input_bytes)
        cmd = [gs, "-sDEVICE=pdfwrite", "-dCompatibilityLevel=1.4", "-dPDFSETTINGS=/ebook", "-dNOPAUSE", "-dQUIET", "-dBATCH", f"-sOutputFile={p_out}", p_in]
        subprocess.run(cmd, timeout=30, check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        if not os.path.exists(p_out):
            return None
        out = Path(p_out).read_bytes()
        if not out or len(out) >= len(input_bytes):
            return None
        try:
            r = pypdf.PdfReader(io.BytesIO(out))
            if len(r.pages) == 0:
                return None
        except Exception:
            return None
        return out
    except Exception:
        return None
    finally:
        for p in (tmp_in, tmp_out):
            if p:
                try:
                    os.remove(p)
                except Exception:
                    pass


def _pypdf_compress(input_bytes):
    try:
        reader = pypdf.PdfReader(io.BytesIO(input_bytes))
        writer = pypdf.PdfWriter()
        for page in reader.pages:
            try:
                page.compress_content_streams(level=9)
            except TypeError:
                try:
                    page.compress_content_streams()
                except Exception:
                    pass
            except Exception:
                pass
            try:
                writer.add_page(page)
            except Exception:
                continue
        # strip document info and XMP metadata
        try:
            if hasattr(writer, "_root_object") and writer._root_object is not None:
                if "/Metadata" in writer._root_object:
                    del writer._root_object["/Metadata"]
        except Exception:
            pass
        # re-encode embedded raster images
        for page in writer.pages:
            try:
                images = list(page.images)
            except Exception:
                continue
            for img_file in images:
                try:
                    orig_data = img_file.data
                    orig_len = len(orig_data) if orig_data else 0
                    pil = img_file.image
                    if pil is None:
                        continue
                    try:
                        pil = ImageOps.exif_transpose(pil)
                    except Exception:
                        pass
                    w, h = pil.size
                    longest = max(w, h)
                    if longest > 2500:
                        ratio = 2500.0 / longest
                        nw, nh = max(1, int(w * ratio)), max(1, int(h * ratio))
                        pil = pil.resize((nw, nh), Image.Resampling.LANCZOS)
                    has_alpha = ("A" in pil.getbands()) if hasattr(pil, "getbands") else (pil.mode in ("RGBA", "LA", "PA"))
                    buf = io.BytesIO()
                    if has_alpha:
                        if pil.mode not in ("RGBA", "LA"):
                            pil = pil.convert("RGBA")
                        pil.save(buf, format="PNG", optimize=True, compress_level=9)
                    else:
                        if pil.mode not in ("RGB", "L"):
                            pil = pil.convert("RGB")
                        pil.save(buf, format="JPEG", quality=75, optimize=True, progressive=True)
                    new_bytes = buf.getvalue()
                    if orig_len and len(new_bytes) >= orig_len:
                        continue
                    if has_alpha:
                        img_file.replace(pil, optimize=True, compress_level=9)
                    else:
                        img_file.replace(pil, quality=75, optimize=True, progressive=True)
                except Exception:
                    continue
        out_buf = io.BytesIO()
        writer.write(out_buf)
        out = out_buf.getvalue()
        if not out:
            return None
        try:
            r = pypdf.PdfReader(io.BytesIO(out))
            if len(r.pages) == 0:
                return None
        except Exception:
            return None
        return out
    except Exception:
        return None


def _render_pdf_pypdfium2(pdf_bytes, dpi, fmt):
    try:
        import pypdfium2
    except Exception:
        return None
    try:
        pdf = pypdfium2.PdfDocument(pdf_bytes)
    except Exception:
        return None
    try:
        n = len(pdf)
        if n == 0:
            return None
        scale = dpi / 72.0
        out = []
        for i in range(n):
            try:
                page = pdf[i]
                pil = page.render(scale=scale).to_pil()
                buf = io.BytesIO()
                if fmt == "JPEG":
                    if pil.mode in ("RGBA", "LA", "P", "PA"):
                        pil = pil.convert("RGB")
                    elif pil.mode != "RGB":
                        try:
                            pil = pil.convert("RGB")
                        except Exception:
                            pil = pil.convert("RGB")
                    pil.save(buf, format="JPEG", quality=85, optimize=True, progressive=True)
                else:
                    pil.save(buf, format="PNG", optimize=True, compress_level=9)
                data = buf.getvalue()
                if not data:
                    return None
                out.append(data)
            except Exception:
                return None
        if not out:
            return None
        return out
    except Exception:
        return None
    finally:
        try:
            pdf.close()
        except Exception:
            pass


def _render_pdf_gs(pdf_bytes, dpi, fmt):
    gs = shutil.which("gs")
    if not gs:
        for cand in ("gswin64c", "gswin32c", "gsc"):
            gs = shutil.which(cand)
            if gs:
                break
    if not gs:
        return None
    tmpdir = None
    tmp_in = None
    try:
        tmpdir = tempfile.mkdtemp()
        fd, tmp_in = tempfile.mkstemp(suffix=".pdf")
        os.close(fd)
        Path(tmp_in).write_bytes(pdf_bytes)
        ext = "png" if fmt == "PNG" else "jpg"
        device = "png16m" if fmt == "PNG" else "jpeg"
        pattern = os.path.join(tmpdir, f"page%d.{ext}")
        if fmt == "PNG":
            cmd = [gs, "-dNOPAUSE", "-dBATCH", f"-sDEVICE={device}", f"-r{dpi}", f"-o{pattern}", tmp_in]
        else:
            cmd = [gs, "-dNOPAUSE", "-dBATCH", f"-sDEVICE={device}", "-dJPEGQ=85", f"-r{dpi}", f"-o{pattern}", tmp_in]
        subprocess.run(cmd, timeout=60, check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        files = []
        for p in Path(tmpdir).glob(f"page*.{ext}"):
            name = p.name
            # extract number between 'page' and '.ext'
            try:
                num = name[len("page"):-len("."+ext)]
                files.append((int(num), p))
            except Exception:
                continue
        if not files:
            return None
        files.sort(key=lambda x: x[0])
        out = []
        for _, p in files:
            try:
                data = p.read_bytes()
                if not data:
                    return None
                out.append(data)
            except Exception:
                return None
        if not out:
            return None
        return out
    except Exception:
        return None
    finally:
        if tmp_in:
            try:
                os.remove(tmp_in)
            except Exception:
                pass
        if tmpdir:
            try:
                shutil.rmtree(tmpdir)
            except Exception:
                pass


def _word_to_pdf_soffice(inp_bytes, filename):
    soff = shutil.which("soffice")
    if not soff:
        return None
    tmpdir = None
    try:
        tmpdir = tempfile.mkdtemp()
        suffix = Path(filename).suffix or ".docx"
        safe = "input" + suffix
        ip = os.path.join(tmpdir, safe)
        Path(ip).write_bytes(inp_bytes)
        cmd = [soff, "--headless", "--convert-to", "pdf", "--outdir", tmpdir, ip]
        subprocess.run(cmd, timeout=120, check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        pdfs = list(Path(tmpdir).glob("*.pdf"))
        if not pdfs:
            exp = Path(tmpdir) / (Path(safe).stem + ".pdf")
            if exp.exists():
                pdfs = [exp]
            else:
                return None
        pp = max(pdfs, key=lambda p: p.stat().st_size) if len(pdfs) > 1 else pdfs[0]
        data = pp.read_bytes()
        if not data or len(data) < 80 or not data.startswith(b"%PDF"):
            return None
        try:
            r = pypdf.PdfReader(io.BytesIO(data))
            if len(r.pages) == 0:
                return None
        except Exception:
            pass
        return data
    except Exception:
        return None
    finally:
        if tmpdir:
            try:
                shutil.rmtree(tmpdir)
            except Exception:
                pass


def _word_to_pdf_docx_pillow(inp_bytes):
    try:
        import docx
    except Exception:
        return None
    try:
        from PIL import Image, ImageDraw, ImageFont
    except Exception:
        return None
    try:
        doc = docx.Document(io.BytesIO(inp_bytes))
    except Exception:
        return None
    W, H = 1240, 1754
    margin = 60
    usable = W - 2 * margin

    def _get_font(sz, bold=False):
        cands = []
        if bold:
            cands += ["DejaVuSans-Bold.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", "C:\\Windows\\Fonts\\DejaVuSans-Bold.ttf", "C:\\Windows\\Fonts\\DejaVuSans.ttf"]
        cands += ["DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "C:\\Windows\\Fonts\\DejaVuSans.ttf", "/system/fonts/DroidSans.ttf"]
        for c in cands:
            try:
                return ImageFont.truetype(c, sz)
            except Exception:
                continue
        try:
            return ImageFont.load_default(size=sz)
        except TypeError:
            return ImageFont.load_default()

    font_normal = _get_font(22)
    font_h1 = _get_font(30, bold=True)
    font_h2 = _get_font(26, bold=True)
    font_h3 = _get_font(23, bold=True)
    font_tbl = _get_font(18)

    def wrap(text, font, draw):
        if not text:
            return []
        words = text.split()
        if not words:
            return [text]
        lines, cur = [], ""
        for w in words:
            tst = cur + (" " if cur else "") + w
            try:
                wd = draw.textlength(tst, font=font)
            except Exception:
                try:
                    wd = font.getlength(tst)
                except Exception:
                    wd = len(tst) * 10
            if wd <= usable:
                cur = tst
            else:
                if cur:
                    lines.append(cur)
                # hard-break long word
                if len(w) > 35:
                    # split word
                    chunk = ""
                    for ch in w:
                        tc = chunk + ch
                        try:
                            cwd = draw.textlength(tc, font=font)
                        except Exception:
                            cwd = len(tc) * 10
                        if cwd > usable:
                            if chunk:
                                lines.append(chunk)
                            chunk = ch
                        else:
                            chunk = tc
                    cur = chunk
                else:
                    cur = w
        if cur:
            lines.append(cur)
        return lines

    pages = []
    im = Image.new("RGB", (W, H), "white")
    draw = ImageDraw.Draw(im)
    y = margin
    line_h_normal = 32
    line_h_h1 = 42
    line_h_tbl = 26

    # map elements to preserve order
    para_map = {p._element: p for p in doc.paragraphs}
    tbl_map = {t._element: t for t in doc.tables}
    try:
        body_children = list(doc.element.body.iterchildren())
    except Exception:
        body_children = list(para_map.keys()) + list(tbl_map.keys())

    def new_page():
        nonlocal im, draw, y
        pages.append(im)
        im = Image.new("RGB", (W, H), "white")
        draw = ImageDraw.Draw(im)
        y = margin

    def ensure(h):
        nonlocal y, im, draw
        if y + h > H - margin:
            new_page()

    for child in body_children:
        try:
            if child in para_map:
                p = para_map[child]
                txt = p.text.strip() if p.text else ""
                if not txt:
                    ensure(line_h_normal // 2)
                    y += line_h_normal // 2
                    continue
                style = ""
                try:
                    style = p.style.name if p.style and p.style.name else ""
                except Exception:
                    style = ""
                if style.startswith("Heading 1") or style == "Title":
                    f, lh = font_h1, line_h_h1
                elif style.startswith("Heading 2"):
                    f, lh = font_h2, 36
                elif style.startswith("Heading 3"):
                    f, lh = font_h3, 33
                else:
                    f, lh = font_normal, line_h_normal
                ls = wrap(txt, f, draw)
                for ln in ls:
                    ensure(lh)
                    try:
                        draw.text((margin, y), ln, fill="black", font=f)
                    except Exception:
                        draw.text((margin, y), ln, fill="black")
                    y += lh
                y += 6
            elif child in tbl_map:
                t = tbl_map[child]
                for row in t.rows:
                    try:
                        cells = []
                        for c in row.cells:
                            try:
                                ct = c.text.strip().replace("\n", " ") if c.text else ""
                            except Exception:
                                ct = ""
                            cells.append(ct)
                        line = " | ".join(cells).strip()
                        if not line:
                            continue
                        ls = wrap(line, font_tbl, draw)
                        for ln in ls:
                            ensure(line_h_tbl)
                            try:
                                draw.text((margin, y), ln, fill="black", font=font_tbl)
                            except Exception:
                                draw.text((margin, y), ln, fill="black")
                            y += line_h_tbl
                    except Exception:
                        continue
                y += 8
            else:
                continue
        except Exception:
            continue

    # fallback if body_children empty (no mapped) then use paragraphs/tables directly
    if not pages and y == margin:
        # try simple paragraphs fallback
        for p in doc.paragraphs:
            try:
                txt = p.text.strip() if p.text else ""
                if not txt:
                    continue
                ls = wrap(txt, font_normal, draw)
                for ln in ls:
                    ensure(line_h_normal)
                    draw.text((margin, y), ln, fill="black", font=font_normal)
                    y += line_h_normal
                y += 4
            except Exception:
                continue
        for t in doc.tables:
            for row in t.rows:
                try:
                    line = " | ".join(c.text.strip().replace("\n", " ") for c in row.cells)
                    if not line.strip():
                        continue
                    ls = wrap(line, font_tbl, draw)
                    for ln in ls:
                        ensure(line_h_tbl)
                        draw.text((margin, y), ln, fill="black", font=font_tbl)
                        y += line_h_tbl
                except Exception:
                    continue

    pages.append(im)
    # validate at least some content: if all blank single page but y==margin -> still count
    try:
        buf = io.BytesIO()
        pages[0].save(buf, "PDF", save_all=True, append_images=pages[1:], resolution=150.0)
        out = buf.getvalue()
        if not out or len(out) < 200:
            return None
        try:
            r = pypdf.PdfReader(io.BytesIO(out))
            if len(r.pages) == 0:
                return None
        except Exception:
            if not out.startswith(b"%PDF"):
                return None
        return out
    except Exception:
        return None


def _pdf_to_word_pdf2docx(inp_bytes):
    try:
        import importlib
        pdf2docx_mod = importlib.import_module("pdf2docx")
        Converter = getattr(pdf2docx_mod, "Converter")
    except Exception:
        return None
    tmp_in = tmp_out = None
    try:
        fd1, p1 = tempfile.mkstemp(suffix=".pdf")
        os.close(fd1)
        fd2, p2 = tempfile.mkstemp(suffix=".docx")
        os.close(fd2)
        tmp_in, tmp_out = p1, p2
        Path(p1).write_bytes(inp_bytes)
        cv = Converter(p1)
        try:
            cv.convert(p2, start=0, end=None)
        finally:
            try:
                cv.close()
            except Exception:
                pass
        data = Path(p2).read_bytes()
        if not data or len(data) < 100:
            return None
        return data
    except Exception:
        return None
    finally:
        for p in (tmp_in, tmp_out):
            if p:
                try:
                    os.remove(p)
                except Exception:
                    pass


def _pdf_to_word_pypdf_docx(inp_bytes):
    try:
        import docx
    except Exception:
        return None
    try:
        reader = pypdf.PdfReader(io.BytesIO(inp_bytes))
        n = len(reader.pages)
        if n == 0:
            return None
    except Exception:
        return None
    try:
        doc = docx.Document()
        for idx, page in enumerate(reader.pages, start=1):
            try:
                txt = page.extract_text() or ""
            except Exception:
                txt = ""
            try:
                doc.add_heading(f"Halaman {idx}", level=1)
            except Exception:
                pa = doc.add_paragraph()
                try:
                    pa.add_run(f"Halaman {idx}").bold = True
                except Exception:
                    pa.add_run(f"Halaman {idx}")
            if txt.strip():
                paras = []
                cur = []
                for line in txt.splitlines():
                    if line.strip() == "":
                        if cur:
                            paras.append("\n".join(cur))
                            cur = []
                    else:
                        cur.append(line)
                if cur:
                    paras.append("\n".join(cur))
                if not paras:
                    paras = [txt.strip()]
                for pt in paras:
                    try:
                        if pt.strip():
                            doc.add_paragraph(pt.strip())
                    except Exception:
                        continue
        buf = io.BytesIO()
        doc.save(buf)
        out = buf.getvalue()
        if not out or len(out) < 80:
            return None
        return out
    except Exception:
        return None


class ToolboxHandler(BaseHTTPRequestHandler):
    def log_message(self, format, *args):
        print(f"[tools] {self.client_address[0]} {format % args}", flush=True)

    def send_json(self, data, code=200):
        b = json.dumps(data).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(b)))
        self.end_headers()
        self.wfile.write(b)

    def send_file_download(self, data_bytes, filename, mime_type="application/octet-stream", original_size=None, page_count=None):
        if original_size is None:
            original_size = len(data_bytes)
        self.send_response(200)
        self.send_header("Content-Type", mime_type)
        self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
        self.send_header("Content-Length", str(len(data_bytes)))
        self.send_header("X-Original-Size", str(int(original_size)))
        self.send_header("X-Result-Size", str(int(len(data_bytes))))
        if page_count is not None:
            self.send_header("X-Page-Count", str(int(page_count)))
        self.end_headers()
        self.wfile.write(data_bytes)

    def do_GET(self):
        parsed = urlparse(self.path)
        path = parsed.path

        if path == "/" or path == "/index.html":
            index_path = STATIC_DIR / "index.html"
            if index_path.exists():
                data = index_path.read_bytes()
                self.send_response(200)
                self.send_header("Content-Type", "text/html; charset=utf-8")
                self.send_header("Content-Length", str(len(data)))
                self.end_headers()
                self.wfile.write(data)
            else:
                self.send_response(404)
                self.end_headers()
            return

        if path == "/api/status":
            self.send_json({"status": "online", "port": PORT, "features": ["pdf_merge", "pdf_compress", "img_to_pdf", "img_compress", "img_upscale", "img_convert", "pdf_to_img", "word_to_pdf", "pdf_to_word"]})
            return

        self.send_response(404)
        self.end_headers()

    def do_POST(self):
        parsed = urlparse(self.path)
        path = parsed.path
        ctype = self.headers.get("Content-Type", "")

        if not ctype.startswith("multipart/form-data"):
            self.send_json({"error": "Expected multipart/form-data"}, 400)
            return

        boundary = ctype.split("boundary=")[-1].encode("utf-8")
        content_length = int(self.headers.get("Content-Length", 0))
        raw_body = self.rfile.read(content_length)

        parts = raw_body.split(b"--" + boundary)
        files = []
        form_data = {}

        for part in parts:
            if not part or part == b"--\r\n" or part == b"--":
                continue
            head_and_body = part.split(b"\r\n\r\n", 1)
            if len(head_and_body) != 2:
                continue
            header_raw, body_raw = head_and_body
            # strip exactly ONE trailing CRLF (the part delimiter) — rstrip would
            # eat real file bytes when uploads legitimately end in CR/LF
            # (e.g. PDFs ending in %%EOF newline -> corrupt X-Original-Size)
            if body_raw.endswith(b"\r\n"):
                body_raw = body_raw[:-2]
            elif body_raw.endswith(b"\n") or body_raw.endswith(b"\r"):
                body_raw = body_raw[:-1]
            header_str = header_raw.decode("utf-8", errors="replace")

            if 'filename="' in header_str:
                fn = header_str.split('filename="')[1].split('"')[0]
                if fn:
                    files.append((fn, body_raw))
            elif 'name="' in header_str:
                field_name = header_str.split('name="')[1].split('"')[0]
                form_data[field_name] = body_raw.decode("utf-8", errors="replace")

        try:
            if path == "/api/pdf/merge":
                if not files:
                    self.send_json({"error": "No PDF files uploaded"}, 400); return
                merger_cls = getattr(pypdf, "PdfMerger", None)
                if merger_cls is not None:
                    merger = merger_cls()
                    for fn, b in files:
                        merger.append(io.BytesIO(b))
                    out_buf = io.BytesIO()
                    merger.write(out_buf)
                    merger.close()
                else:
                    writer = pypdf.PdfWriter()
                    for fn, b in files:
                        writer.append(io.BytesIO(b))
                    out_buf = io.BytesIO()
                    writer.write(out_buf)
                    writer.close()
                out_bytes = out_buf.getvalue()
                orig = sum(len(b) for _, b in files)
                self.send_file_download(out_bytes, "merged_document.pdf", "application/pdf", original_size=orig)
                return

            elif path == "/api/pdf/compress":
                if not files:
                    self.send_json({"error": "No PDF file uploaded"}, 400); return
                fn, b = files[0]
                orig_len = len(b)
                best = b
                gs_bytes = None
                try:
                    gs_bytes = _gs_compress(b)
                except Exception:
                    gs_bytes = None
                if gs_bytes and len(gs_bytes) < len(best):
                    best = gs_bytes
                pypdf_bytes = None
                try:
                    pypdf_bytes = _pypdf_compress(b)
                except Exception:
                    pypdf_bytes = None
                if pypdf_bytes and len(pypdf_bytes) < len(best):
                    best = pypdf_bytes
                if len(best) > orig_len:
                    best = b
                self.send_file_download(best, f"compressed_{fn}", "application/pdf", original_size=orig_len)
                return

            elif path == "/api/pdf/to-img":
                if not files:
                    self.send_json({"error": "No PDF file uploaded"}, 400); return
                fn, b = files[0]
                orig_len = len(b)
                fmt = form_data.get("format", "PNG").upper().strip()
                if fmt == "JPG":
                    fmt = "JPEG"
                if fmt not in ("PNG", "JPEG"):
                    fmt = "PNG"
                try:
                    dpi = int(form_data.get("dpi", "150"))
                except Exception:
                    dpi = 150
                if dpi not in (96, 150, 300):
                    dpi = 150
                stem = Path(fn).stem if Path(fn).stem else "document"
                # quick validate PDF header to give clear error for non-PDF
                try:
                    pypdf.PdfReader(io.BytesIO(b))
                except Exception:
                    # let render layer try, but if both fail will return engine or invalid error
                    # we still attempt render; if fails we return invalid PDF error
                    pass
                images = None
                try:
                    images = _render_pdf_pypdfium2(b, dpi, fmt)
                except Exception:
                    images = None
                if not images:
                    try:
                        images = _render_pdf_gs(b, dpi, fmt)
                    except Exception:
                        images = None
                if not images or len(images) == 0 or any(not x for x in images):
                    # check if any engine exists
                    has_pdfium = False
                    try:
                        import pypdfium2
                        has_pdfium = True
                    except Exception:
                        has_pdfium = False
                    has_gs = shutil.which("gs") is not None
                    if not has_gs:
                        for cand in ("gswin64c", "gswin32c", "gsc"):
                            if shutil.which(cand):
                                has_gs = True
                                break
                    if not has_pdfium and not has_gs:
                        self.send_json({"error": "Render engine tidak tersedia. Install: pip install pypdfium2 (atau pkg install ghostscript di Termux)"}, 500); return
                    self.send_json({"error": "Failed to render PDF"}, 500); return
                page_count = len(images)
                # validate images are openable
                try:
                    for data in images:
                        im = Image.open(io.BytesIO(data))
                        im.verify()
                except Exception:
                    self.send_json({"error": "Failed to render PDF"}, 500); return
                ext = "png" if fmt == "PNG" else "jpg"
                if page_count == 1:
                    data = images[0]
                    mime = "image/png" if fmt == "PNG" else "image/jpeg"
                    filename = f"{stem}_page1.{ext}"
                    self.send_response(200)
                    self.send_header("Content-Type", mime)
                    self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
                    self.send_header("Content-Length", str(len(data)))
                    self.send_header("X-Original-Size", str(int(orig_len)))
                    self.send_header("X-Result-Size", str(int(len(data))))
                    self.send_header("X-Page-Count", str(int(page_count)))
                    self.end_headers()
                    self.wfile.write(data)
                    return
                else:
                    zip_buf = io.BytesIO()
                    with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
                        for idx, data in enumerate(images, start=1):
                            arcname = f"{stem}_page{idx}.{ext}"
                            zf.writestr(arcname, data)
                    zip_bytes = zip_buf.getvalue()
                    filename = f"fotopdf_{stem}.zip"
                    self.send_response(200)
                    self.send_header("Content-Type", "application/zip")
                    self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
                    self.send_header("Content-Length", str(len(zip_bytes)))
                    self.send_header("X-Original-Size", str(int(orig_len)))
                    self.send_header("X-Result-Size", str(int(len(zip_bytes))))
                    self.send_header("X-Page-Count", str(int(page_count)))
                    self.end_headers()
                    self.wfile.write(zip_bytes)
                    return

            elif path == "/api/img/to-pdf":
                if not files:
                    self.send_json({"error": "No image files uploaded"}, 400); return
                pil_imgs = []
                for fn, b in files:
                    try:
                        im = Image.open(io.BytesIO(b))
                        try:
                            im = ImageOps.exif_transpose(im)
                        except Exception:
                            pass
                        if im.mode in ("RGBA", "P"):
                            im = im.convert("RGB")
                        pil_imgs.append(im)
                    except Exception:
                        pass
                if not pil_imgs:
                    self.send_json({"error": "Invalid image files"}, 400); return
                out_buf = io.BytesIO()
                pil_imgs[0].save(out_buf, "PDF", save_all=True, append_images=pil_imgs[1:], resolution=100.0)
                out_bytes = out_buf.getvalue()
                orig = sum(len(b) for _, b in files)
                self.send_file_download(out_bytes, "images_document.pdf", "application/pdf", original_size=orig)
                return

            elif path == "/api/img/compress":
                if not files:
                    self.send_json({"error": "No image uploaded"}, 400); return
                fn, b = files[0]
                orig_len = len(b)
                quality = int(form_data.get("quality", 70))
                quality = max(1, min(95, quality))
                im = Image.open(io.BytesIO(b))
                # capture fmt/icc BEFORE exif_transpose: transpose returns a new
                # image object whose .format is None (would fall back to JPEG)
                fmt = im.format if im.format in ("JPEG", "PNG", "WEBP") else "JPEG"
                icc = im.info.get("icc_profile")
                try:
                    im = ImageOps.exif_transpose(im)
                except Exception:
                    pass
                if fmt == "JPEG" and im.mode in ("RGBA", "P"):
                    im = im.convert("RGB")
                if fmt == "WEBP" and im.mode == "P":
                    try:
                        im = im.convert("RGBA")
                    except Exception:
                        im = im.convert("RGB")
                out_buf = io.BytesIO()
                try:
                    if fmt == "JPEG":
                        if icc:
                            im.save(out_buf, format="JPEG", quality=quality, optimize=True, progressive=True, icc_profile=icc)
                        else:
                            im.save(out_buf, format="JPEG", quality=quality, optimize=True, progressive=True)
                    elif fmt == "PNG":
                        tmp1 = io.BytesIO()
                        if icc:
                            im.save(tmp1, format="PNG", optimize=True, compress_level=9, icc_profile=icc)
                        else:
                            im.save(tmp1, format="PNG", optimize=True, compress_level=9)
                        best_bytes = tmp1.getvalue()
                        best_len = len(best_bytes)
                        try:
                            has_alpha = ("A" in im.getbands()) if hasattr(im, "getbands") else (im.mode == "RGBA")
                            if has_alpha:
                                rgba = im.convert("RGBA") if im.mode != "RGBA" else im
                                q_im = rgba.quantize(colors=256, method=Image.Quantize.FASTOCTREE)
                            else:
                                rgb = im.convert("RGB") if im.mode != "RGB" else im
                                q_im = rgb.quantize(colors=256, method=Image.Quantize.MEDIANCUT)
                            if q_im is not None:
                                tmp2 = io.BytesIO()
                                q_im.save(tmp2, format="PNG", optimize=True, compress_level=9)
                                q_bytes = tmp2.getvalue()
                                if len(q_bytes) < best_len:
                                    best_bytes = q_bytes
                        except Exception:
                            pass
                        out_buf = io.BytesIO(best_bytes)
                    elif fmt == "WEBP":
                        if icc:
                            im.save(out_buf, format="WEBP", quality=quality, method=6, icc_profile=icc)
                        else:
                            im.save(out_buf, format="WEBP", quality=quality, method=6)
                    else:
                        im.save(out_buf, format=fmt, quality=quality, optimize=True)
                except Exception:
                    out_buf = io.BytesIO()
                    try:
                        im.save(out_buf, format=fmt, quality=quality, optimize=True)
                    except Exception:
                        out_buf = io.BytesIO(b)
                result = out_buf.getvalue()
                if not result or len(result) > orig_len:
                    result = b
                ext = "jpg" if fmt == "JPEG" else fmt.lower()
                mime = "image/jpeg" if fmt == "JPEG" else f"image/{ext}"
                if result is b:
                    orig_ext = Path(fn).suffix.lstrip(".").lower() or ext
                    if orig_ext in ("jpg", "jpeg"):
                        ext2 = "jpg"
                        mime = "image/jpeg"
                    else:
                        ext2 = orig_ext
                        mime = f"image/{ext2}" if ext2 in ("png", "webp") else f"image/{ext}"
                    self.send_file_download(result, f"compressed_{Path(fn).stem}.{ext2}", mime, original_size=orig_len)
                else:
                    self.send_file_download(result, f"compressed_{Path(fn).stem}.{ext}", mime, original_size=orig_len)
                return

            elif path == "/api/img/upscale":
                if not files:
                    self.send_json({"error": "No image uploaded"}, 400); return
                fn, b = files[0]
                orig_len = len(b)
                scale = int(form_data.get("scale", 2))
                if scale not in (2, 4): scale = 2
                im = Image.open(io.BytesIO(b))
                # capture fmt/icc BEFORE exif_transpose (new object loses .format)
                fmt = im.format if im.format in ("JPEG", "PNG", "WEBP") else "PNG"
                icc = im.info.get("icc_profile")
                try:
                    im = ImageOps.exif_transpose(im)
                except Exception:
                    pass
                new_w = im.width * scale
                new_h = im.height * scale
                upscaled = im.resize((new_w, new_h), Image.Resampling.LANCZOS)
                enhancer = ImageEnhance.Sharpness(upscaled)
                upscaled = enhancer.enhance(1.25)
                out_buf = io.BytesIO()
                if fmt == "JPEG" and upscaled.mode in ("RGBA", "P"):
                    upscaled = upscaled.convert("RGB")
                try:
                    if fmt == "JPEG":
                        if icc:
                            upscaled.save(out_buf, format="JPEG", quality=95, optimize=True, progressive=True, icc_profile=icc)
                        else:
                            upscaled.save(out_buf, format="JPEG", quality=95, optimize=True, progressive=True)
                    elif fmt == "PNG":
                        if icc:
                            upscaled.save(out_buf, format="PNG", optimize=True, compress_level=9, icc_profile=icc)
                        else:
                            upscaled.save(out_buf, format="PNG", optimize=True, compress_level=9)
                    elif fmt == "WEBP":
                        if icc:
                            upscaled.save(out_buf, format="WEBP", quality=95, method=6, icc_profile=icc)
                        else:
                            upscaled.save(out_buf, format="WEBP", quality=95, method=6)
                    else:
                        upscaled.save(out_buf, format=fmt, quality=95, optimize=True)
                except Exception:
                    upscaled.save(out_buf, format=fmt, quality=95)
                ext = "jpg" if fmt == "JPEG" else fmt.lower()
                self.send_file_download(out_buf.getvalue(), f"upscaled_{scale}x_{Path(fn).stem}.{ext}", f"image/{ext}", original_size=orig_len)
                return

            elif path == "/api/img/convert":
                if not files:
                    self.send_json({"error": "No image uploaded"}, 400); return
                fn, b = files[0]
                orig_len = len(b)
                target_fmt = form_data.get("format", "PNG").upper()
                if target_fmt not in ("PNG", "JPEG", "JPG", "WEBP"):
                    target_fmt = "PNG"
                if target_fmt == "JPG":
                    target_fmt = "JPEG"
                im = Image.open(io.BytesIO(b))
                try:
                    im = ImageOps.exif_transpose(im)
                except Exception:
                    pass
                icc = im.info.get("icc_profile")
                out_buf = io.BytesIO()
                if target_fmt == "JPEG" and im.mode in ("RGBA", "P"):
                    im = im.convert("RGB")
                try:
                    if target_fmt == "JPEG":
                        if icc:
                            im.save(out_buf, format="JPEG", quality=90, optimize=True, progressive=True, icc_profile=icc)
                        else:
                            im.save(out_buf, format="JPEG", quality=90, optimize=True, progressive=True)
                    elif target_fmt == "PNG":
                        if icc:
                            im.save(out_buf, format="PNG", optimize=True, compress_level=9, icc_profile=icc)
                        else:
                            im.save(out_buf, format="PNG", optimize=True, compress_level=9)
                    elif target_fmt == "WEBP":
                        if icc:
                            im.save(out_buf, format="WEBP", quality=90, method=6, icc_profile=icc)
                        else:
                            im.save(out_buf, format="WEBP", quality=90, method=6)
                    else:
                        im.save(out_buf, format=target_fmt, quality=90)
                except Exception:
                    im.save(out_buf, format=target_fmt, quality=90)
                ext = "jpg" if target_fmt == "JPEG" else target_fmt.lower()
                self.send_file_download(out_buf.getvalue(), f"converted_{Path(fn).stem}.{ext}", f"image/{ext}", original_size=orig_len)
                return

            elif path == "/api/word/to-pdf":
                if not files:
                    self.send_json({"error": "No file uploaded"}, 400); return
                fn, b = files[0]
                orig_len = len(b)
                stem = Path(fn).stem or "document"
                pdf_bytes = _word_to_pdf_soffice(b, fn)
                if not pdf_bytes:
                    is_legacy = fn.lower().endswith(".doc") or b[:4] == b'\xd0\xcf\x11\xe0'
                    if is_legacy and not zipfile.is_zipfile(io.BytesIO(b)):
                        self.send_json({"error": ".doc butuh LibreOffice (pkg install libreoffice)"}, 400); return
                    pdf_bytes = _word_to_pdf_docx_pillow(b)
                    if not pdf_bytes:
                        self.send_json({"error": "Konversi Word tidak tersedia. Install: pip install python-docx (atau libreoffice buat hasil penuh)"}, 500); return
                self.send_file_download(pdf_bytes, f"converted_{stem}.pdf", "application/pdf", original_size=orig_len)
                return

            elif path == "/api/pdf/to-word":
                if not files:
                    self.send_json({"error": "No file uploaded"}, 400); return
                fn, b = files[0]
                orig_len = len(b)
                stem = Path(fn).stem or "document"
                docx_bytes = _pdf_to_word_pdf2docx(b)
                if not docx_bytes:
                    docx_bytes = _pdf_to_word_pypdf_docx(b)
                    if not docx_bytes:
                        self.send_json({"error": "Konversi PDF ke Word tidak tersedia. Install: pip install pdf2docx / python-docx"}, 500); return
                self.send_file_download(docx_bytes, f"converted_{stem}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", original_size=orig_len)
                return

            self.send_json({"error": "Unknown action"}, 404)
        except Exception as e:
            self.send_json({"error": str(e)}, 500)

if __name__ == "__main__":
    print(f"Gwolf Toolbox listening on http://127.0.0.1:{PORT}")
    server = ThreadingHTTPServer(("127.0.0.1", PORT), ToolboxHandler)
    server.serve_forever()
